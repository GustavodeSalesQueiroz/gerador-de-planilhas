import json
from io import BytesIO

from django.test import Client, TestCase
from django.urls import reverse
from openpyxl import load_workbook

from .models import Planilha, RegistroConversa, TemaBot


class ChatbotTests(TestCase):
    def setUp(self):
        self.event_theme = TemaBot.objects.get(codigo="evento")
        self.event_answers = {
            "nome_evento": "Feira Geek",
            "data_evento": "2026-11-04",
            "local": "Centro cultural",
            "status": "Planejado",
        }

    def post_conversation(self, theme="evento", answers=None):
        return self.client.post(
            reverse("criar_conversa"),
            data=json.dumps({"tema": theme, "respostas": answers or self.event_answers}),
            content_type="application/json",
        )

    def create_sheet(self, name="Feira de novembro", answers=None):
        return self.client.post(
            reverse("criar_planilha"),
            data=json.dumps({
                "tema": "evento",
                "nome": name,
                "respostas": answers or self.event_answers,
            }),
            content_type="application/json",
        )

    def test_home_shows_chatbot_and_recent_records(self):
        response = self.client.get(reverse("home"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Assistente de registros")
        self.assertContains(response, "chatbot-config")

    def test_script_returns_questions_messages_and_event_theme_from_database(self):
        response = self.client.get(reverse("roteiro_chatbot"))

        self.assertEqual(response.status_code, 200)
        data = response.json()
        event = next(theme for theme in data["temas"] if theme["codigo"] == "evento")
        self.assertEqual(event["nome"], "Eventos")
        self.assertEqual(event["perguntas"][0]["pergunta"], "Qual é o nome do evento?")
        self.assertEqual(data["mensagens"]["confirmar"], "Confirmar e criar planilha")
        self.assertEqual({theme["codigo"] for theme in data["temas"]}, {
            "pedido", "venda", "produto", "estoque", "evento",
        })

    def test_script_reflects_question_edited_in_database(self):
        question = self.event_theme.perguntas.get(campo="nome_evento")
        question.pergunta = "Como se chama o evento?"
        question.save()

        response = self.client.get(reverse("roteiro_chatbot"))
        event = next(theme for theme in response.json()["temas"] if theme["codigo"] == "evento")

        self.assertEqual(event["perguntas"][0]["pergunta"], "Como se chama o evento?")

    def test_saves_valid_conversation_and_returns_download_link(self):
        response = self.post_conversation()

        self.assertEqual(response.status_code, 201)
        self.assertEqual(RegistroConversa.objects.count(), 1)
        conversa = RegistroConversa.objects.get()
        self.assertEqual(conversa.tema, self.event_theme)
        self.assertEqual(conversa.respostas["nome_evento"], "Feira Geek")
        self.assertEqual(response.json()["download_url"], "/planilhas/evento/baixar/")

    def test_rejects_missing_answer(self):
        answers = self.event_answers.copy()
        answers["local"] = ""

        response = self.post_conversation(answers=answers)

        self.assertEqual(response.status_code, 400)
        self.assertIn("Local precisa ser preenchido", response.json()["erro"])
        self.assertEqual(RegistroConversa.objects.count(), 0)

    def test_rejects_invalid_date_number_and_choice(self):
        invalid_answers = [
            {**self.event_answers, "data_evento": "not-a-date"},
            {**self.event_answers, "status": "Não existe"},
        ]
        sales = TemaBot.objects.get(codigo="venda")
        sales_answers = {
            "produto": "Miniatura",
            "quantidade": "-1",
            "valor_total": "12.00",
            "data_venda": "2026-11-04",
        }

        for answers in invalid_answers:
            with self.subTest(answers=answers):
                response = self.post_conversation(answers=answers)
                self.assertEqual(response.status_code, 400)

        response = self.post_conversation(theme=sales.codigo, answers=sales_answers)
        self.assertEqual(response.status_code, 400)
        self.assertEqual(RegistroConversa.objects.count(), 0)

    def test_rejects_extra_answer_fields_and_malformed_json(self):
        answers = {**self.event_answers, "unexpected": "value"}
        response = self.post_conversation(answers=answers)
        self.assertEqual(response.status_code, 400)

        response = self.client.post(
            reverse("criar_conversa"),
            data="{",
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 400)
        self.assertEqual(RegistroConversa.objects.count(), 0)

    def test_rejects_inactive_or_unknown_theme_with_json_error(self):
        response = self.post_conversation(theme="tema-inexistente")

        self.assertEqual(response.status_code, 404)
        self.assertIn("Este tema não está disponível", response.json()["erro"])
        self.assertEqual(RegistroConversa.objects.count(), 0)

    def test_downloads_excel_file_with_typed_columns_and_theme_records(self):
        self.post_conversation()
        self.post_conversation(answers={
            **self.event_answers,
            "nome_evento": "=HYPERLINK(\"https://example.com\")",
        })

        response = self.client.get(reverse("baixar_planilha", args=["evento"]))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn(".xlsx", response["Content-Disposition"])
        workbook = load_workbook(BytesIO(b"".join(response.streaming_content)), data_only=False)
        sheet = workbook["Eventos"]
        self.assertEqual(sheet.max_row, 3)
        self.assertEqual(sheet["B1"].value, "Nome do evento")
        self.assertEqual(sheet["B2"].value, "Feira Geek")
        self.assertEqual(sheet["B3"].data_type, "s")
        self.assertEqual(sheet["C2"].value.date().isoformat(), "2026-11-04")
        self.assertEqual(sheet["A1"].value, "Data do registro")

    def test_download_contains_only_records_for_selected_theme(self):
        self.post_conversation()
        products = TemaBot.objects.get(codigo="produto")
        RegistroConversa.objects.create(
            tema=products,
            respostas={"nome_produto": "Caneca", "descricao": "Edição especial", "preco": "15.00"},
        )

        response = self.client.get(reverse("baixar_planilha", args=["evento"]))
        workbook = load_workbook(BytesIO(b"".join(response.streaming_content)))

        self.assertEqual(workbook.active.title, "Eventos")
        self.assertEqual(workbook.active.max_row, 2)

    def test_sales_workbook_contains_numeric_values(self):
        answers = {
            "produto": "Miniatura",
            "quantidade": "2",
            "valor_total": "45.50",
            "data_venda": "2026-11-04",
        }
        self.post_conversation(theme="venda", answers=answers)

        response = self.client.get(reverse("baixar_planilha", args=["venda"]))
        workbook = load_workbook(BytesIO(b"".join(response.streaming_content)))
        sheet = workbook["Vendas"]

        self.assertEqual(sheet["C2"].value, 2)
        self.assertEqual(sheet["C2"].data_type, "n")
        self.assertEqual(sheet["D2"].value, 45.5)
        self.assertEqual(sheet["D2"].data_type, "n")

    def test_rejects_request_without_csrf_token(self):
        client = Client(enforce_csrf_checks=True)
        client.get(reverse("home"))

        response = client.post(
            reverse("criar_conversa"),
            data=json.dumps({"tema": "evento", "respostas": self.event_answers}),
            content_type="application/json",
        )

        self.assertEqual(response.status_code, 403)
        self.assertEqual(RegistroConversa.objects.count(), 0)

    def test_creates_a_named_spreadsheet_with_first_record(self):
        response = self.create_sheet()

        self.assertEqual(response.status_code, 201)
        sheet = Planilha.objects.get()
        self.assertEqual(sheet.nome, "Feira de novembro")
        self.assertEqual(sheet.registros.count(), 1)
        self.assertEqual(response.json()["planilha"]["quantidade"], 1)
        self.assertEqual(
            response.json()["download_url"],
            reverse("baixar_planilha_por_id", args=[sheet.pk]),
        )

    def test_adds_multiple_records_to_one_spreadsheet_and_exports_them_together(self):
        created = self.create_sheet()
        sheet_id = created.json()["planilha"]["id"]
        second_answers = {
            **self.event_answers,
            "nome_evento": "Encontro de colecionadores",
            "local": "Praça central",
        }

        response = self.client.post(
            reverse("adicionar_registro_planilha", args=[sheet_id]),
            data=json.dumps({"respostas": second_answers}),
            content_type="application/json",
        )

        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.json()["planilha"]["quantidade"], 2)
        self.assertEqual(Planilha.objects.get(pk=sheet_id).registros.count(), 2)
        download = self.client.get(reverse("baixar_planilha_por_id", args=[sheet_id]))
        workbook = load_workbook(BytesIO(b"".join(download.streaming_content)))
        self.assertEqual(workbook.active.max_row, 3)
        self.assertEqual(workbook.active["B2"].value, "Feira Geek")
        self.assertEqual(workbook.active["B3"].value, "Encontro de colecionadores")

    def test_edits_an_existing_record_and_download_contains_updated_value(self):
        created = self.create_sheet()
        sheet_id = created.json()["planilha"]["id"]
        record_id = created.json()["registro"]["id"]
        updated_answers = {
            **self.event_answers,
            "nome_evento": "Feira Geek Atualizada",
            "local": "Novo local",
        }

        response = self.client.patch(
            reverse("editar_registro_planilha", args=[sheet_id, record_id]),
            data=json.dumps({"respostas": updated_answers}),
            content_type="application/json",
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["mensagem"], "Registro atualizado com sucesso!")
        self.assertEqual(response.json()["planilha"]["quantidade"], 1)
        self.assertEqual(
            Planilha.objects.get(pk=sheet_id).registros.get(pk=record_id).respostas["nome_evento"],
            "Feira Geek Atualizada",
        )
        download = self.client.get(reverse("baixar_planilha_por_id", args=[sheet_id]))
        workbook = load_workbook(BytesIO(b"".join(download.streaming_content)))
        self.assertEqual(workbook.active["B2"].value, "Feira Geek Atualizada")

    def test_separate_spreadsheets_keep_records_isolated(self):
        first = self.create_sheet("Primeiro evento")
        second = self.create_sheet(
            "Segundo evento",
            answers={**self.event_answers, "nome_evento": "Outro evento"},
        )

        self.assertNotEqual(first.json()["planilha"]["id"], second.json()["planilha"]["id"])
        first_download = self.client.get(reverse(
            "baixar_planilha_por_id",
            args=[first.json()["planilha"]["id"]],
        ))
        workbook = load_workbook(BytesIO(b"".join(first_download.streaming_content)))
        self.assertEqual(workbook.active.max_row, 2)
        self.assertEqual(workbook.active["B2"].value, "Feira Geek")

    def test_existing_spreadsheet_list_and_detail_include_rows(self):
        created = self.create_sheet()
        sheet_id = created.json()["planilha"]["id"]

        listing = self.client.get(reverse("listar_planilhas"))
        detail = self.client.get(reverse("detalhe_planilha", args=[sheet_id]))

        self.assertEqual(listing.status_code, 200)
        self.assertEqual(listing.json()["planilhas"][0]["nome"], "Feira de novembro")
        self.assertEqual(detail.status_code, 200)
        self.assertEqual(len(detail.json()["registros"]), 1)
        self.assertEqual(detail.json()["perguntas"][0]["campo"], "nome_evento")

    def test_rejects_invalid_spreadsheet_name_and_invalid_update(self):
        invalid_name = self.create_sheet("   ")
        too_long_name = self.create_sheet("x" * 101)

        self.assertEqual(invalid_name.status_code, 400)
        self.assertEqual(too_long_name.status_code, 400)
        self.assertEqual(Planilha.objects.count(), 0)

        created = self.create_sheet()
        response = self.client.patch(
            reverse(
                "editar_registro_planilha",
                args=[created.json()["planilha"]["id"], created.json()["registro"]["id"]],
            ),
            data=json.dumps({"respostas": {**self.event_answers, "status": "Desconhecido"}}),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 400)
        self.assertEqual(
            Planilha.objects.get().registros.get().respostas["status"],
            "Planejado",
        )

    def test_script_and_conversation_endpoints_only_accept_expected_methods(self):
        self.assertEqual(self.client.post(reverse("roteiro_chatbot")).status_code, 405)
        self.assertEqual(self.client.get(reverse("criar_conversa")).status_code, 405)
        self.assertEqual(self.client.get(reverse("criar_planilha")).status_code, 405)

    def test_collection_and_record_endpoints_reject_wrong_write_method(self):
        created = self.create_sheet()
        sheet_id = created.json()["planilha"]["id"]
        record_id = created.json()["registro"]["id"]

        patch_collection = self.client.patch(
            reverse("adicionar_registro_planilha", args=[sheet_id]),
            data=json.dumps({"respostas": self.event_answers}),
            content_type="application/json",
        )
        post_record = self.client.post(
            reverse("editar_registro_planilha", args=[sheet_id, record_id]),
            data=json.dumps({"respostas": self.event_answers}),
            content_type="application/json",
        )

        self.assertEqual(patch_collection.status_code, 405)
        self.assertEqual(post_record.status_code, 405)
