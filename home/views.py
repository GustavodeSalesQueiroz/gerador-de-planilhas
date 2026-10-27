import json
import re
from datetime import datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.db import transaction
from django.db.models import Count
from django.http import FileResponse, JsonResponse
from django.shortcuts import get_object_or_404, render
from django.template.defaultfilters import slugify
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_GET, require_http_methods, require_POST
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import MensagemBot, Planilha, RegistroConversa, RegistroOperacional, TemaBot


def home(request):
    planilhas = Planilha.objects.select_related("tema").annotate(
        total_registros=Count("registros"),
    )[:10]
    historico = [
        {
            "planilha_id": planilha.pk,
            "tema": planilha.tema.nome,
            "resumo": f"{planilha.nome} · {planilha.total_registros} registro(s)",
            "criado_em": timezone.localtime(planilha.atualizada_em),
        }
        for planilha in planilhas
    ]

    registros_antigos = RegistroOperacional.objects.order_by("-criado_em")[:10]
    for registro in registros_antigos:
        historico.append({
            "tema": registro.get_categoria_display(),
            "resumo": registro.detalhes,
            "criado_em": timezone.localtime(registro.criado_em),
        })
    historico.sort(key=lambda item: item["criado_em"], reverse=True)

    return render(
        request,
        "home.html",
        {
            "registros": historico[:10],
            "chatbot_config": {
                "roteiro_url": reverse("roteiro_chatbot"),
                "planilhas_url": reverse("listar_planilhas"),
                "criar_planilha_url": reverse("criar_planilha"),
            },
        },
    )


@require_GET
def roteiro_chatbot(request):
    temas = TemaBot.objects.filter(ativo=True).prefetch_related("perguntas")
    mensagens = dict(MensagemBot.objects.values_list("chave", "texto"))
    roteiro = [
        {
            "codigo": tema.codigo,
            "nome": tema.nome,
            "perguntas": [
                {
                    "campo": pergunta.campo,
                    "cabecalho": pergunta.cabecalho,
                    "pergunta": pergunta.pergunta,
                    "tipo": pergunta.tipo,
                    "obrigatoria": pergunta.obrigatoria,
                    "opcoes": pergunta.opcoes,
                }
                for pergunta in tema.perguntas.all()
            ],
        }
        for tema in temas
    ]
    return JsonResponse({"mensagens": mensagens, "temas": roteiro})


def texto_bot(chave):
    return MensagemBot.objects.get(chave=chave).texto


def ler_json(request):
    try:
        dados = json.loads(request.body)
    except (json.JSONDecodeError, UnicodeDecodeError):
        return None, JsonResponse({"erro": texto_bot("erro_formato")}, status=400)
    if not isinstance(dados, dict):
        return None, JsonResponse({"erro": texto_bot("erro_formato")}, status=400)
    return dados, None


def validar_respostas(tema, respostas):
    perguntas = list(tema.perguntas.all())
    campos_esperados = {pergunta.campo for pergunta in perguntas}
    if set(respostas) != campos_esperados:
        return None, texto_bot("erro_respostas")

    respostas_validadas = {}
    for pergunta in perguntas:
        valor = respostas.get(pergunta.campo)
        if not isinstance(valor, str):
            return None, texto_bot("erro_campo_obrigatorio").format(campo=pergunta.cabecalho)
        valor = valor.strip()

        if pergunta.obrigatoria and not valor:
            return None, texto_bot("erro_campo_obrigatorio").format(campo=pergunta.cabecalho)
        if not valor:
            respostas_validadas[pergunta.campo] = ""
            continue
        if len(valor) > 300:
            return None, texto_bot("erro_tamanho").format(campo=pergunta.cabecalho)
        if pergunta.opcoes and valor not in pergunta.opcoes:
            return None, texto_bot("erro_opcao").format(campo=pergunta.cabecalho.lower())

        if pergunta.tipo == "number":
            try:
                numero = Decimal(valor)
            except InvalidOperation:
                return None, texto_bot("erro_numero").format(campo=pergunta.cabecalho.lower())
            if not numero.is_finite() or numero < 0 or numero > Decimal("999999999.99"):
                return None, texto_bot("erro_intervalo").format(campo=pergunta.cabecalho.lower())
            valor = format(numero, "f")
        elif pergunta.tipo == "date":
            try:
                datetime.strptime(valor, "%Y-%m-%d")
            except ValueError:
                return None, texto_bot("erro_data").format(campo=pergunta.cabecalho.lower())

        respostas_validadas[pergunta.campo] = valor

    return respostas_validadas, None


def resumo_registro(tema, respostas):
    perguntas = {pergunta.campo: pergunta for pergunta in tema.perguntas.all()}
    return " · ".join(
        f"{perguntas[campo].cabecalho}: {valor}"
        for campo, valor in respostas.items()
        if campo in perguntas
    )


def serializar_registro(registro):
    return {
        "id": registro.pk,
        "planilha_id": registro.planilha_id,
        "respostas": registro.respostas,
        "resumo": resumo_registro(registro.tema, registro.respostas),
        "criado_em": timezone.localtime(registro.criado_em).strftime("%d/%m/%Y %H:%M"),
    }


def serializar_planilha(planilha):
    return {
        "id": planilha.pk,
        "nome": planilha.nome,
        "tema": planilha.tema.codigo,
        "tema_nome": planilha.tema.nome,
        "quantidade": planilha.registros.count(),
        "atualizada_em": timezone.localtime(planilha.atualizada_em).strftime("%d/%m/%Y %H:%M"),
        "url": reverse("baixar_planilha_por_id", args=[planilha.pk]),
    }


def resposta_salva(planilha, registro, mensagem_chave, status=200):
    planilha.atualizada_em = timezone.now()
    planilha.save(update_fields=["atualizada_em"])
    return JsonResponse(
        {
            "mensagem": texto_bot(mensagem_chave),
            "mensagem_download": texto_bot("download_planilha"),
            "mensagem_novo": texto_bot("iniciar_outro"),
            "mensagem_adicionar": texto_bot("adicionar_registro"),
            "mensagem_editar": texto_bot("editar_planilha"),
            "mensagem_editar_linha": texto_bot("editar_registro"),
            "planilha": serializar_planilha(planilha),
            "registro": serializar_registro(registro),
            "download_url": reverse("baixar_planilha_por_id", args=[planilha.pk]),
        },
        status=status,
    )


@require_GET
def listar_planilhas(request):
    planilhas = Planilha.objects.select_related("tema").annotate(
        total_registros=Count("registros"),
    )
    return JsonResponse({
        "planilhas": [
            {
                **serializar_planilha(planilha),
                "quantidade": planilha.total_registros,
            }
            for planilha in planilhas
        ],
    })


@require_GET
def detalhe_planilha(request, planilha_id):
    planilha = get_object_or_404(
        Planilha.objects.select_related("tema").prefetch_related("registros"),
        pk=planilha_id,
    )
    return JsonResponse({
        "planilha": serializar_planilha(planilha),
        "perguntas": [
            {
                "campo": pergunta.campo,
                "cabecalho": pergunta.cabecalho,
                "pergunta": pergunta.pergunta,
                "tipo": pergunta.tipo,
                "obrigatoria": pergunta.obrigatoria,
                "opcoes": pergunta.opcoes,
            }
            for pergunta in planilha.tema.perguntas.all()
        ],
        "registros": [
            serializar_registro(registro)
            for registro in planilha.registros.all()
        ],
    })


@require_POST
def criar_planilha(request):
    dados, erro_json = ler_json(request)
    if erro_json:
        return erro_json

    codigo_tema = dados.get("tema")
    nome = dados.get("nome")
    respostas = dados.get("respostas")
    if not isinstance(codigo_tema, str) or not isinstance(respostas, dict):
        return JsonResponse({"erro": texto_bot("erro_respostas")}, status=400)
    if not isinstance(nome, str) or not nome.strip() or len(nome.strip()) > 100:
        return JsonResponse({"erro": texto_bot("erro_nome_planilha")}, status=400)

    tema = TemaBot.objects.filter(codigo=codigo_tema, ativo=True).first()
    if tema is None:
        return JsonResponse({"erro": texto_bot("erro_tema")}, status=404)
    respostas_validadas, erro = validar_respostas(tema, respostas)
    if erro:
        return JsonResponse({"erro": erro}, status=400)

    with transaction.atomic():
        planilha = Planilha.objects.create(tema=tema, nome=nome.strip())
        registro = RegistroConversa.objects.create(
            tema=tema,
            planilha=planilha,
            respostas=respostas_validadas,
        )
    return resposta_salva(planilha, registro, "planilha_criada", status=201)


@require_http_methods(["POST", "PATCH"])
def salvar_registro_planilha(request, planilha_id, registro_id=None):
    planilha = get_object_or_404(
        Planilha.objects.select_related("tema"),
        pk=planilha_id,
    )
    dados, erro_json = ler_json(request)
    if erro_json:
        return erro_json

    respostas = dados.get("respostas")
    if not isinstance(respostas, dict):
        return JsonResponse({"erro": texto_bot("erro_respostas")}, status=400)
    respostas_validadas, erro = validar_respostas(planilha.tema, respostas)
    if erro:
        return JsonResponse({"erro": erro}, status=400)

    if request.method == "POST":
        if registro_id is not None:
            return JsonResponse(
                {"erro": texto_bot("erro_respostas")},
                status=405,
                headers={"Allow": "PATCH"},
            )
        registro = RegistroConversa.objects.create(
            tema=planilha.tema,
            planilha=planilha,
            respostas=respostas_validadas,
        )
        return resposta_salva(planilha, registro, "registro_salvo", status=201)

    if registro_id is None:
        return JsonResponse(
            {"erro": texto_bot("erro_respostas")},
            status=405,
            headers={"Allow": "POST"},
        )

    registro = get_object_or_404(
        RegistroConversa,
        pk=registro_id,
        planilha=planilha,
    )
    registro.respostas = respostas_validadas
    registro.save(update_fields=["respostas"])
    return resposta_salva(planilha, registro, "registro_atualizado")


@require_POST
def criar_conversa(request):
    dados, erro_json = ler_json(request)
    if erro_json:
        return erro_json

    codigo_tema = dados.get("tema")
    respostas = dados.get("respostas")
    if not isinstance(codigo_tema, str) or not isinstance(respostas, dict):
        return JsonResponse({"erro": texto_bot("erro_respostas")}, status=400)

    tema = TemaBot.objects.filter(codigo=codigo_tema, ativo=True).first()
    if tema is None:
        return JsonResponse({"erro": texto_bot("erro_tema")}, status=404)
    respostas_validadas, erro = validar_respostas(tema, respostas)
    if erro:
        return JsonResponse({"erro": erro}, status=400)

    conversa = RegistroConversa.objects.create(tema=tema, respostas=respostas_validadas)
    return JsonResponse(
        {
            "mensagem": texto_bot("registro_salvo"),
            "mensagem_download": texto_bot("download_planilha"),
            "mensagem_novo": texto_bot("iniciar_outro"),
            "registro": {
                "tema": tema.nome,
                "resumo": resumo_registro(tema, respostas_validadas),
                "criado_em": timezone.localtime(conversa.criado_em).strftime("%d/%m/%Y %H:%M"),
            },
            "download_url": reverse("baixar_planilha", args=[tema.codigo]),
        },
        status=201,
    )


def gerar_arquivo_excel(planilha, registros, nome_arquivo):
    tema = planilha.tema
    perguntas = list(tema.perguntas.all())
    workbook = Workbook()
    sheet = workbook.active
    titulo = re.sub(r"[\[\]:*?/\\]", " ", planilha.nome).strip("'")[:31]
    sheet.title = titulo or tema.nome_planilha
    sheet.append(["Data do registro"] + [pergunta.cabecalho for pergunta in perguntas])

    for registro in registros:
        created_at = timezone.localtime(registro.criado_em).replace(tzinfo=None)
        row = [created_at]
        for pergunta in perguntas:
            valor = registro.respostas.get(pergunta.campo, "")
            if pergunta.tipo == "number" and valor:
                valor = float(Decimal(valor))
            elif pergunta.tipo == "date" and valor:
                valor = datetime.strptime(valor, "%Y-%m-%d").date()
            row.append(valor)
        sheet.append(row)
        for column_index, pergunta in enumerate(perguntas, start=2):
            if pergunta.tipo == "text":
                sheet.cell(row=sheet.max_row, column=column_index).data_type = "s"

    header_fill = PatternFill(fill_type="solid", fgColor="245C4A")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.data_type = "s"

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        max_width = max(len(str(cell.value or "")) for cell in column)
        sheet.column_dimensions[get_column_letter(column[0].column)].width = min(max(max_width + 2, 14), 40)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return FileResponse(
        output,
        as_attachment=True,
        filename=nome_arquivo,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@require_GET
def baixar_planilha_por_id(request, planilha_id):
    planilha = get_object_or_404(
        Planilha.objects.select_related("tema").prefetch_related("registros"),
        pk=planilha_id,
    )
    nome = slugify(planilha.nome) or planilha.tema.codigo
    filename = f"{nome}-{planilha.pk}.xlsx"
    return gerar_arquivo_excel(
        planilha,
        planilha.registros.order_by("criado_em", "pk"),
        filename,
    )


@require_GET
def baixar_planilha(request, codigo):
    tema = get_object_or_404(TemaBot, codigo=codigo, ativo=True)
    planilha = Planilha(tema=tema, nome=tema.nome_planilha)
    registros = RegistroConversa.objects.filter(tema=tema, planilha__isnull=True)
    filename = f"{tema.codigo}-{timezone.localtime():%Y%m%d-%H%M%S}.xlsx"
    return gerar_arquivo_excel(planilha, registros.order_by("criado_em", "pk"), filename)
